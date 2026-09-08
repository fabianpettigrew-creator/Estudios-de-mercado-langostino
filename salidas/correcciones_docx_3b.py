# -*- coding: utf-8 -*-
"""Ajustes de maqueta del Anexo A: parte la ecuación (A.10) en dos renglones, evita
títulos huérfanos y hace que el encabezado de las tablas se repita al cortar página."""
from docx import Document
from docx.shared import Pt, RGBColor, Twips
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

RUTA = "/tmp/lango/out/Demanda_inversa_langostino_metodologia_y_resultados.docx"
TXT, GRIS = "1A1A1A", "5A5A5A"
CENTRO, DERECHA = 4550, 9100
doc = Document(RUTA)

# ---- 1. (A.10) en dos renglones -------------------------------------------------
obj = None
for p in doc.paragraphs:
    if p.text.strip().startswith("w_tan,t = α + γ · log(q_tan,t / q_fre,t)"):
        obj = p
if obj is None:
    raise SystemExit("no se encontró (A.10)")
for r in list(obj.runs):
    r._r.getparent().remove(r._r)
L1 = "w_tan,t  =  α + γ · log(q_tan,t / q_fre,t) + β · log Q*_t + θ · log p*_t"
L2 = "+ φ·h_t + τ·t + Σ_m δ_m·D_mt + ε_t"
r = obj.add_run("\t" + L1)
r.font.bold = True; r.font.color.rgb = RGBColor.from_string(TXT)
r.add_break()
r2 = obj.add_run("\t" + L2)
r2.font.bold = True; r2.font.color.rgb = RGBColor.from_string(TXT)
r3 = obj.add_run("\t(A.10)")
r3.font.color.rgb = RGBColor.from_string(GRIS); r3.font.size = Pt(9)
print("(A.10) partida en dos renglones")

# ---- 2. títulos que no se separan de lo que sigue --------------------------------
inicio = None
for i, p in enumerate(doc.paragraphs):
    if p.text.startswith("Anexo A."):
        inicio = i; break
n = 0
for p in doc.paragraphs[inicio:]:
    if p.style.name.startswith("Heading"):
        p.paragraph_format.keep_with_next = True; n += 1
    # el párrafo que presenta una tabla tampoco debe quedar solo al pie
    if p.text.rstrip().endswith(":") or p.text.rstrip().endswith("es:"):
        p.paragraph_format.keep_with_next = True
print(f"keep_with_next en {n} títulos del anexo")

# ---- 3. encabezado repetido en las tablas del anexo ------------------------------
for t in doc.tables[-4:]:
    trPr = t.rows[0]._tr.get_or_add_trPr()
    if trPr.find(qn("w:tblHeader")) is None:
        e = OxmlElement("w:tblHeader"); e.set(qn("w:val"), "true"); trPr.append(e)
    # y ninguna fila se parte por la mitad
    for fila in t.rows:
        pr = fila._tr.get_or_add_trPr()
        if pr.find(qn("w:cantSplit")) is None:
            c = OxmlElement("w:cantSplit"); pr.append(c)
print("encabezados repetidos y filas indivisibles en las 4 tablas del anexo")

doc.save(RUTA)
print("guardado")
