# -*- coding: utf-8 -*-
"""Precio mensual del langostino argentino L1 puesto en España, contra el camarón
ecuatoriano puesto en España. Enero de 2020 a julio de 2026, toneladas, valor y US$/kg.

Las dos puntas salen de despachos de exportación transacción por transacción (base de
comercio), así que son precios FOB del país de origen, comparables entre sí.

UN LÍMITE QUE HAY QUE LEER ANTES DE USAR LA PLANILLA. Del lado argentino el calibre L1
está identificado en la descripción comercial de cada despacho. Del lado ecuatoriano NO
se puede filtrar por talla: la descripción declara talla en el 22% de las filas en 2020
y se desploma a 0,6% en 2025, de modo que no hay serie mensual emparejada por calibre.
Lo que va en la columna de Ecuador es TODO el camarón de la partida 0306.17.99
despachado a España, sin filtro de talla, que es lo máximo que el dato permite.
La hoja «Talla en Ecuador» documenta esa cobertura año por año.

Y una segunda advertencia: el flujo ecuatoriano a España es mayormente intra-grupo
—Promarisco, brazo ecuatoriano de Pescanova, embarcando a Pescanova España—, así que
sus valores unitarios son en parte precios de transferencia y no precios de mercado.

Uso:  python l1_espana_ar_ec.py
Salida: salidas/L1_Espana_Argentina_vs_Ecuador_2020_2026.xlsx
"""
from __future__ import annotations

import glob
import os
import sys

import numpy as np
import openpyxl
import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L

sys.stdout.reconfigure(encoding="utf-8")
RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, RAIZ)
EXPO_EC = os.environ.get("EXPO_ECUADOR_DIR", os.path.join(RAIZ, "..", "EXPO ECUADOR"))
SALIDA = os.path.join(RAIZ, "salidas",
                      "L1_Espana_Argentina_vs_Ecuador_2020_2026.xlsx")
DESDE, HASTA = "2020-01", "2026-07"


# ------------------------------------------------------------------ Argentina
def argentina() -> pd.DataFrame:
    """Entero calibre L1 despachado a España, por mes y por flota."""
    from fuentes import expo
    e = expo.cargar_todos()
    d = e[(e.pres == "entero") & (e.cal == "L1") & (e.kg > 0) & (e.fob_tot > 0)
          & e.destino.astype(str).str.upper().str.strip()
          .isin(["ESPAÑA", "ESPANA"])].copy()
    d["flota"] = np.where(d.ruta == "a bordo", "tangonera", "fresquera")
    d["per"] = pd.PeriodIndex(d.anio.astype(str) + "-"
                              + d.mes.astype(str).str.zfill(2), freq="M")
    return d[(d.per >= pd.Period(DESDE)) & (d.per <= pd.Period(HASTA))]


# -------------------------------------------------------------------- Ecuador
def ecuador() -> pd.DataFrame:
    """Camarón de la partida 0306.17.99 despachado a España, transacción por
    transacción. Sin filtro de talla, porque el dato no lo permite."""
    fs = sorted(glob.glob(os.path.join(EXPO_EC, "Softrade_EC_export_*Espana*.xlsx")))
    if not fs:
        sys.exit(f"ABORTA: no encontré los despachos de Ecuador en {EXPO_EC}")
    S = pd.concat([pd.read_excel(f, sheet_name="Exportaciones") for f in fs],
                  ignore_index=True)
    S.columns = [str(c).strip() for c in S.columns]
    S["Fecha"] = pd.to_datetime(S["Fecha"], errors="coerce")
    S["kg"] = pd.to_numeric(S["Kgs. Netos"], errors="coerce")
    S["fob"] = pd.to_numeric(S["U$S FOB"], errors="coerce")
    S = S[(S.kg > 0) & (S.fob > 0) & S.Fecha.notna()].copy()
    S["per"] = S.Fecha.dt.to_period("M")
    S["anio"] = S.Fecha.dt.year
    d = S["Descripción Comercial"].astype(str).str.upper()
    t = d.str.extract(r"(?:T\.?\s*|TALLA\s*)(\d{1,3})\s*[-/]\s*(\d{1,3})")
    S["t0"] = pd.to_numeric(t[0], errors="coerce")
    S["t1"] = pd.to_numeric(t[1], errors="coerce")
    S["exportador"] = S.Exportador.astype(str).str.strip()
    return S[(S.per >= pd.Period(DESDE)) & (S.per <= pd.Period(HASTA))]


def _mes(g: pd.DataFrame, kg="kg", fob="fob") -> pd.DataFrame:
    x = g.groupby("per").agg(t=(kg, "sum"), fob=(fob, "sum"), n=(kg, "size"))
    x["t"] /= 1000.0
    x["usd_kg"] = x.fob / (x.t * 1000.0)
    return x


# ------------------------------------------------------------------- estilos
AZUL, AZUL2, GRIS = "1F4E79", "2E6DA4", "F2F5F8"
AR = Font(name="Arial", size=10)
ARB = Font(name="Arial", size=10, bold=True)
HDR = Font(name="Arial", size=10, bold=True, color="FFFFFF")
TIT = Font(name="Arial", size=13, bold=True, color=AZUL)
SUB = Font(name="Arial", size=9, color="5A5A5A")
AZB = Font(name="Arial", size=10, bold=True, color=AZUL)
ROJO = Font(name="Arial", size=9, bold=True, color="C0504D")
FA, FA2 = PatternFill("solid", fgColor=AZUL), PatternFill("solid", fgColor=AZUL2)
FG = PatternFill("solid", fgColor=GRIS)
CEN = Alignment(horizontal="center", vertical="center", wrap_text=True)
DER = Alignment(horizontal="right")
SEP = Border(left=Side(style="thin", color="BFBFBF"))


def encabezado(ws, titulo, bajada, aviso=None):
    ws["A1"] = titulo
    ws["A1"].font = TIT
    ws["A2"] = bajada
    ws["A2"].font = SUB
    ws["A2"].alignment = Alignment(vertical="top", wrap_text=True)
    if aviso:
        ws["A3"] = aviso
        ws["A3"].font = ROJO


def tabla(ws, r0, grupos, idx, filas, fmts):
    """`grupos` son (título, ncols); `filas` un dict per → lista de valores."""
    j = 2
    for tit, n in grupos:
        ws.merge_cells(start_row=r0, start_column=j, end_row=r0, end_column=j + n - 1)
        c = ws.cell(r0, j, tit)
        c.font, c.fill, c.alignment = HDR, FA, CEN
        j += n
    r1 = r0 + 1
    enc = ["Mes"] + [e for _, _, e in filas]
    for j, t in enumerate(enc, start=1):
        c = ws.cell(r1, j, t)
        c.font, c.fill, c.alignment = HDR, FA2, CEN
    ws.row_dimensions[r1].height = 30

    r = r1 + 1
    cortes = set()
    acc = 2
    for _, n in grupos:
        cortes.add(acc)
        acc += n
    for per in idx:
        gris = FG if per.year % 2 else None
        ws.cell(r, 1, f"{per.year}-{per.month:02d}").font = AR
        for j, ((serie, fmt, _), ) in enumerate(zip([(a, b, c) for a, b, c in filas]),
                                                start=2):
            v = serie.get(per, np.nan)
            c = ws.cell(r, j, None if pd.isna(v) else float(v))
            c.font, c.number_format, c.alignment = AR, fmt, DER
            if j in cortes:
                c.border = SEP
        if gris:
            for j in range(1, len(filas) + 2):
                ws.cell(r, j).fill = gris
        r += 1
    ws.column_dimensions["A"].width = 10
    for j in range(2, len(filas) + 2):
        ws.column_dimensions[L(j)].width = 13
    ws.freeze_panes = "B" + str(r1 + 1)
    return r


def main() -> None:
    A = argentina()
    E = ecuador()
    idx = pd.period_range(DESDE, HASTA, freq="M")
    mA, mE = _mes(A, "kg", "fob_tot"), _mes(E)
    mT = _mes(A[A.flota == "tangonera"], "kg", "fob_tot")
    mF = _mes(A[A.flota == "fresquera"], "kg", "fob_tot")
    print(f"Argentina L1 entero a España: {len(A):,} despachos · {mA.t.sum():,.0f} t")
    print(f"Ecuador 0306.17.99 a España:  {len(E):,} despachos · {mE.t.sum():,.0f} t")

    coc = (mA.usd_kg / mE.usd_kg).reindex(idx)
    dif = ((mA.usd_kg / mE.usd_kg - 1) * 100).reindex(idx)

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    BAJ = ("Despachos de exportación transacción por transacción (base de comercio). "
           "Argentina: langostino ENTERO calibre L1, partida 0306.17, destino España. "
           "Ecuador: camarón\nde la partida 0306.17.99, destino España, SIN filtro de "
           "talla. Toneladas de producto, FOB declarado en dólares y precio ponderado "
           "por volumen.")
    AVISO = ("ATENCIÓN: las dos columnas no son el mismo producto. Ver la hoja «Notas» "
             "antes de citar el cociente.")

    # -------------------------------------------------- 1. comparativa
    ws = wb.create_sheet("Comparativa mensual")
    encabezado(ws, "Langostino argentino L1 contra camarón ecuatoriano, puestos en "
                   "España — 2020 a julio 2026", BAJ, AVISO)
    tabla(ws, 5,
          [("ARGENTINA · entero L1", 3), ("ECUADOR · 0306.17.99, toda talla", 3),
           ("COMPARACIÓN", 2)],
          idx,
          [(mA.t, "#,##0.0", "Toneladas"),
           (mA.fob, "#,##0", "FOB (US$)"),
           (mA.usd_kg, "#,##0.00", "US$/kg"),
           (mE.t, "#,##0.0", "Toneladas"),
           (mE.fob, "#,##0", "FOB (US$)"),
           (mE.usd_kg, "#,##0.00", "US$/kg"),
           (coc, "#,##0.000", "Cociente AR / EC"),
           (dif, '+#,##0.0"%";-#,##0.0"%"', "Prima argentina")],
          None)

    # -------------------------------------------------- 2. Argentina por flota
    ws = wb.create_sheet("Argentina por flota")
    encabezado(ws, "Langostino argentino entero L1 a España, por flota — 2020 a julio 2026",
               "Flota tangonera = congelado a bordo; fresquera = procesado en tierra. "
               "Despachos de exportación (base de comercio).")
    tabla(ws, 4,
          [("TANGONERA", 3), ("FRESQUERA", 3), ("TOTAL", 3)],
          idx,
          [(mT.t, "#,##0.0", "Toneladas"), (mT.fob, "#,##0", "FOB (US$)"),
           (mT.usd_kg, "#,##0.00", "US$/kg"),
           (mF.t, "#,##0.0", "Toneladas"), (mF.fob, "#,##0", "FOB (US$)"),
           (mF.usd_kg, "#,##0.00", "US$/kg"),
           (mA.t, "#,##0.0", "Toneladas"), (mA.fob, "#,##0", "FOB (US$)"),
           (mA.usd_kg, "#,##0.00", "US$/kg")],
          None)

    # -------------------------------------------------- 3. resumen anual
    ws = wb.create_sheet("Resumen anual")
    encabezado(ws, "Resumen anual — 2020 a julio 2026", BAJ)
    enc = ["Año", "AR · toneladas", "AR · FOB (millones US$)", "AR · US$/kg",
           "EC · toneladas", "EC · FOB (millones US$)", "EC · US$/kg",
           "Cociente AR / EC", "Prima argentina", "AR · despachos", "EC · despachos"]
    for j, t in enumerate(enc, start=1):
        c = ws.cell(4, j, t)
        c.font, c.fill, c.alignment = HDR, FA2, CEN
    ws.row_dimensions[4].height = 30
    aA = A.groupby(A.per.dt.year).agg(t=("kg", "sum"), f=("fob_tot", "sum"),
                                      n=("kg", "size"))
    aE = E.groupby(E.per.dt.year).agg(t=("kg", "sum"), f=("fob", "sum"),
                                      n=("kg", "size"))
    r = 5
    for i, y in enumerate(sorted(aA.index)):
        pa, pe = aA.f[y] / aA.t[y], aE.f[y] / aE.t[y]
        vals = [y, aA.t[y] / 1000, aA.f[y] / 1e6, pa, aE.t[y] / 1000, aE.f[y] / 1e6,
                pe, pa / pe, (pa / pe - 1) * 100, int(aA.n[y]), int(aE.n[y])]
        fmts = ["0", "#,##0", "#,##0.0", "#,##0.00", "#,##0", "#,##0.0", "#,##0.00",
                "#,##0.000", '+#,##0.0"%";-#,##0.0"%"', "#,##0", "#,##0"]
        for j, (v, fm) in enumerate(zip(vals, fmts), start=1):
            c = ws.cell(r, j, v)
            c.font, c.number_format, c.alignment = AR, fm, DER
            if i % 2:
                c.fill = FG
        r += 1
    for j, w in enumerate([8, 14, 17, 12, 14, 17, 12, 14, 13, 13, 13], start=1):
        ws.column_dimensions[L(j)].width = w
    ws.freeze_panes = "B5"

    # -------------------------------------------------- 4. diagnóstico de talla
    ws = wb.create_sheet("Talla en Ecuador")
    encabezado(ws, "Por qué la columna de Ecuador no está filtrada por calibre",
               "Cobertura de la talla en la descripción comercial de cada despacho "
               "ecuatoriano a España, y concentración del exportador.")
    enc = ["Año", "Despachos", "Toneladas", "US$/kg", "% de filas con talla",
           "% de kilos con talla", "Meses con dato de talla", "Exportador principal",
           "% de kilos del principal"]
    for j, t in enumerate(enc, start=1):
        c = ws.cell(4, j, t)
        c.font, c.fill, c.alignment = HDR, FA2, CEN
    ws.row_dimensions[4].height = 30
    r = 5
    for i, (y, s) in enumerate(E.groupby("anio")):
        ex = s.groupby("exportador").kg.sum().sort_values(ascending=False)
        vals = [int(y), len(s), s.kg.sum() / 1000, s.fob.sum() / s.kg.sum(),
                s.t0.notna().mean() * 100,
                s.loc[s.t0.notna(), "kg"].sum() / s.kg.sum() * 100,
                int(s.loc[s.t0.notna(), "per"].nunique()),
                ex.index[0], ex.iloc[0] / s.kg.sum() * 100]
        fmts = ["0", "#,##0", "#,##0.0", "#,##0.00", "#,##0.0", "#,##0.0", "0", None,
                "#,##0.0"]
        for j, (v, fm) in enumerate(zip(vals, fmts), start=1):
            c = ws.cell(r, j, v)
            c.font = AR
            if fm:
                c.number_format, c.alignment = fm, DER
            if i % 2:
                c.fill = FG
        r += 1
    for j, w in enumerate([8, 12, 12, 10, 16, 16, 16, 30, 16], start=1):
        ws.column_dimensions[L(j)].width = w

    r += 2
    ws.cell(r, 1, "Las bandas de talla que sí aparecen, en todo el período").font = AZB
    r += 1
    for j, t in enumerate(["Banda", "Despachos", "Toneladas", "US$/kg"], start=1):
        c = ws.cell(r, j, t)
        c.font, c.fill, c.alignment = HDR, FA2, CEN
    B = E[E.t0.notna()].copy()
    B["banda"] = B.t0.astype(int).astype(str) + "/" + B.t1.astype(int).astype(str)
    g = B.groupby("banda").agg(n=("kg", "size"), kg=("kg", "sum"), f=("fob", "sum"))
    g["t"] = g.kg / 1000
    g["usd_kg"] = g.f / g.kg
    for i, (b, f) in enumerate(g.sort_values("kg", ascending=False).head(12).iterrows()):
        r += 1
        for j, (v, fm) in enumerate(zip([b, int(f.n), f.t, f.usd_kg],
                                        [None, "#,##0", "#,##0.0", "#,##0.00"]), start=1):
            c = ws.cell(r, j, v)
            c.font = AR
            if fm:
                c.number_format, c.alignment = fm, DER
            if i % 2:
                c.fill = FG
    r += 2
    ws.cell(r, 1, "El L1 argentino son 11 a 20 piezas por kilo. Ninguna banda "
                  "ecuatoriana con volumen relevante llega a ese tamaño, y las que "
                  "aparecen suman menos del 3% de los kilos del período.").font = SUB

    # -------------------------------------------------- 5. notas
    NOTAS = [
        ("Lo primero: las dos columnas NO son el mismo producto", True),
        ("Del lado argentino la columna es langostino entero de calibre L1 —11 a 20 "
         "piezas por kilo—, identificado en la descripción comercial de cada despacho. "
         "Del lado ecuatoriano es TODO el camarón de la partida 0306.17.99 despachado a "
         "España, de cualquier talla y cualquier presentación.", False),
        ("No se puede hacer mejor con este dato. La descripción comercial ecuatoriana "
         "declara talla en el 21,9% de las filas en 2020 y se desploma a 0,6% en 2025; "
         "medido en kilos la cobertura va de 12,3% a 0,1%. Con eso no hay serie mensual "
         "emparejada por calibre: hay meses enteros sin una sola fila con talla. La hoja "
         "«Talla en Ecuador» tiene el detalle año por año.", False),
        ("Consecuencia práctica: el cociente y la prima de la hoja comparativa mezclan "
         "diferencia de PRECIO con diferencia de TAMAÑO. Sirven para seguir el "
         "movimiento en el tiempo, no para afirmar cuánto más caro es el langostino que "
         "el vannamei a igual calibre.", False),
        ("", False),
        ("Lo segundo: el flujo ecuatoriano a España es mayormente intra-grupo", True),
        ("Promarisco S.A., brazo ecuatoriano de Pescanova, embarca la mayor parte de "
         "los kilos, y Pescanova España los recibe. Los valores unitarios de un flujo "
         "así son en parte precios de transferencia y no precios de mercado. La "
         "columna «% de kilos del principal» de la hoja «Talla en Ecuador» muestra la "
         "concentración año por año.", False),
        ("", False),
        ("Qué miden las cifras", True),
        ("Las toneladas son de PRODUCTO tal cual se despacha, no peso de langostino "
         "entero equivalente. El valor es FOB declarado en dólares. El precio es el FOB "
         "sobre los kilos, ponderado por volumen dentro de cada mes: no es una "
         "cotización, y se mueve también con la mezcla de producto de cada embarque.", False),
        ("Las dos puntas son FOB del país de origen, así que son comparables entre sí. "
         "No son comparables con los precios CIF de importación de la UE que publica "
         "EUMOFA, que incluyen flete y seguro hasta el puerto europeo.", False),
        ("", False),
        ("Cobertura", True),
        ("Enero de 2020 a julio de 2026 en las dos puntas. Del lado argentino la "
         "ventana arranca en 2020 porque es desde donde el volumen del despacho es "
         "confiable: antes hay doble conteo que infla los kilos hasta 1,45 veces, "
         "aunque el precio no se ve afectado porque es un cociente.", False),
        ("Un mes en blanco significa que no hubo despachos, no que falte el dato.", False),
        ("", False),
        ("Reproducibilidad", True),
        ("«salidas/l1_espana_ar_ec.py». Elaboración: Lic. Fabián Pettigrew · AXIA.", False),
    ]
    ws = wb.create_sheet("Notas")
    ws.column_dimensions["A"].width = 118
    for i, (t, neg) in enumerate(NOTAS, start=1):
        c = ws.cell(i, 1, t)
        c.font = Font(name="Arial", size=10, bold=neg, color=AZUL if neg else "1A1A1A")
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if not neg and t:
            ws.row_dimensions[i].height = 14 * (1 + len(t) // 108)

    for h in wb.sheetnames:
        wb[h].sheet_properties.tabColor = "7F7F7F" if h in ("Notas", "Talla en Ecuador") else AZUL
    wb.save(SALIDA)
    print("escrito:", SALIDA)

    print("\nanual:")
    for y in sorted(aA.index):
        pa, pe = aA.f[y] / aA.t[y], aE.f[y] / aE.t[y]
        print(f"   {y}   AR {aA.t[y] / 1000:7,.0f} t  {pa:5.2f} US$/kg   |   "
              f"EC {aE.t[y] / 1000:7,.0f} t  {pe:5.2f} US$/kg   |   "
              f"prima {(pa / pe - 1) * 100:+5.1f}%")


if __name__ == "__main__":
    main()
