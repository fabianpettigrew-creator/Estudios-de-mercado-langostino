# -*- coding: utf-8 -*-
"""Saca del .docx de metodología el eje del COSTO EVITABLE / MARGEN y la viñeta del
exportador. Decisión de Fabián, 24-ago-2026. Se parchea sobre el archivo vigente."""
from docx import Document

SRC = '/tmp/lango/base2.docx'
OUT = '/tmp/lango/out/Demanda_inversa_langostino_metodologia_y_resultados.docx'
d = Document(SRC)


def buscar(frag):
    for p in d.paragraphs:
        if frag in p.text:
            return p
    raise SystemExit(f'NO ENCONTRADO: {frag!r}')


def borrar(frag):
    p = buscar(frag)
    p._element.getparent().remove(p._element)
    print(f'   borrado: {frag[:60]}…')


def renombrar(viejo, nuevo):
    p = buscar(viejo)
    p.runs[0].text = nuevo
    for r in p.runs[1:]:
        r.text = ''
    print(f'   «{viejo}» → «{nuevo}»')


print('1. §1 resumen ejecutivo')
borrar('Pero el test que decide no es el de la facturación sino el del margen')

print('2. §2 el test del margen')
borrar('El test del margen. A la empresa no le importa la facturación')

print('3. Cuadro 12 — se cae la columna del costo evitable')
tb = None
for t in d.tables:
    if t.rows and 'Costo evitable' in t.rows[0].cells[-1].text:
        tb = t
        break
assert tb is not None, 'no encontré el Cuadro 12'
for fila in tb.rows:
    celda = fila.cells[-1]._tc
    celda.getparent().remove(celda)
grid = tb._tbl.tblGrid
grid.remove(grid[-1])
print(f'   Cuadro 12 quedó en {len(tb.columns)} columnas')

print('4. epígrafe y cierre del §10')
borrar('La última columna es el costo evitable por kilo')
p = buscar('Sobre los US$867 millones de exportación')
p.runs[0].text = ('Sobre los US$867 millones de exportación de langostino de 2025, la estimación '
                  'preferida implica una caída de facturación del orden de US$52 millones por año.')
for r in p.runs[1:]:
    r.text = ''
print('   cierre del §10 reescrito sin el umbral de margen')

print('5. §11.1 completo')
borrar('11.1 Lo que puede cambiar la recomendación')
borrar('Un solo dato: la estructura de costo evitable del buque tangonero')

print('6. viñeta del exportador')
borrar('La razón social del exportador, hoy no disponible en ninguna fila')

print('7. renumeración de los apartados del §11')
renombrar('11.2 Lo que afina el número', '11.1 Lo que afina el número')
renombrar('11.3 Limitaciones declaradas', '11.2 Limitaciones declaradas')

d.save(OUT)
print('\nguardado en', OUT)
