# -*- coding: utf-8 -*-
"""¿Cuánto de la divergencia de precio entre el langostino argentino y el vannamei en
España es precio y cuánto es cambio de mezcla de producto?

EL PROBLEMA QUE VIENE A RESOLVER. La elasticidad de sustitución se estima con valores
unitarios: valor declarado sobre kilos. Un valor unitario sube por dos motivos
distintos, que el dato de Comext no separa porque la nomenclatura a 8 dígitos no tiene
calibre: porque subió el precio del mismo producto, o porque se embarcó producto más
caro. Si la oferta argentina a España se corrió hacia producto de mayor valor a lo
largo de la serie, parte de lo que el modelo de sustitución atribuye al precio es
composición, y la elasticidad está sesgada.

CÓMO SE ACOTA. La base de comercio sí trae calibre y presentación despacho por
despacho, en el campo AI de la descripción comercial. Con eso se arma la canasta
argentina a España y se calcula un contrafáctico: qué valor unitario habría tenido cada
año si la mezcla se hubiera quedado congelada en la de 2013, valuada a los precios de
cada año. La diferencia entre el valor unitario observado y ese contrafáctico es el
efecto composición.

    ln(composición) = ln( valor unitario observado / valor unitario con el mix de 2013 )

Las categorías se agrupan en entero L1, entero L2, entero L3, cola y pelado. Más fino
que eso no se sostiene: dentro de cola el grado no se declara de manera consistente.
La canasta base de 2013 no tiene pelado, así que esa categoría entra con peso cero en
el contrafáctico, que es exactamente lo que corresponde: su aparición ES composición.

QUÉ SE HACE CON EL RESULTADO. Se corrige el precio relativo de Comext sacándole el
factor de composición y se vuelve a estimar la elasticidad con la misma especificación
DOLS. Si el coeficiente se mueve poco, la composición no estaba manejando el resultado.

LÍMITES QUE NO SE PUEDEN SALVAR CON ESTE DATO.
  * La cobertura de calibre va del 96% de los kilos en 2024-2026 al 36% en 2017. El
    índice se calcula sobre el subconjunto que declara calibre y se supone que el
    resto se comporta igual.
  * La base de comercio es FOB en Argentina y Comext es CIF en frontera española. El
    factor de composición se mide en una punta y se aplica en la otra. La correlación
    entre las variaciones anuales de los dos valores unitarios es 0,39: co-mueven,
    pero no son la misma serie. La corrección es aproximada y así hay que citarla.
  * Antes de 2020 el volumen de los despachos tiene doble conteo. Las
    participaciones son cocientes, así que el sesgo se cancela si la duplicación no
    discrimina por tipo de producto. No hay manera de verificarlo.

Uso:  python composicion_calibre_es.py
Salida: consola
        salidas/Composicion_calibre_Espana.xlsx
        salidas/espana_composicion_calibre.svg  (+ .png)
"""
from __future__ import annotations

import os
import sys
import warnings

import numpy as np
import openpyxl
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter as L

warnings.filterwarnings("ignore")
sys.stdout.reconfigure(encoding="utf-8")
RAIZ = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, RAIZ)
SAL = os.path.join(RAIZ, "salidas")

import elasticidad_sustitucion_es as E   # noqa: E402
from fuentes import expo                 # noqa: E402

BASE = 2013                              # año de la canasta congelada
AZUL, ROJO, GRIS = "#1f4e79", "#c0504d", "#9aa7ad"
COLORES = {"entero L1": "#1f4e79", "entero L2": "#4a7ebb", "entero L3": "#a3c1e0",
           "cola": "#c0504d", "pelado": "#e8a87c"}


# ---------------------------------------------------------------------- datos
def despachos() -> pd.DataFrame:
    """Despachos argentinos a España con calibre y presentación identificados."""
    d = expo.cargar_todos()
    e = d[d.destino.astype(str).str.upper().str.strip().str.startswith("ESPA")].copy()
    cob = (e.loc[e.cal != "s/d", "kg"].sum() / e.kg.sum())
    print(f"Despachos a España: {len(e):,} · {e.kg.sum() / 1e6:,.1f} kt · "
          f"calibre identificado en {100 * cob:.1f}% de los kilos")
    e = e[(e.kg > 0) & (e.fob_tot > 0) & (e.cal != "s/d") & (e.pres != "otro")].copy()
    e["cat"] = np.where(e.pres == "entero", "entero " + e.cal,
                        np.where(e.pres == "cola", "cola", "pelado"))
    e["per"] = pd.PeriodIndex(e.anio.astype(str) + "-"
                              + e.mes.astype(str).str.zfill(2), freq="M")
    return e


def cobertura(e_todo: pd.DataFrame) -> pd.DataFrame:
    d = expo.cargar_todos()
    t = d[d.destino.astype(str).str.upper().str.strip().str.startswith("ESPA")]
    f = t.groupby("anio").apply(lambda x: pd.Series({
        "Despachos": len(x), "Kilotoneladas": x.kg.sum() / 1e6,
        "% de kilos con calibre": 100 * x.loc[x.cal != "s/d", "kg"].sum() / x.kg.sum(),
        "% de kilos con presentación":
            100 * x.loc[x.pres != "otro", "kg"].sum() / x.kg.sum()}))
    return f.reset_index().rename(columns={"anio": "Año"})


def indice(e: pd.DataFrame, freq: str):
    """Valor unitario observado y contrafáctico con la canasta congelada en BASE."""
    key = e.per if freq == "M" else e.anio
    g = e.groupby([key, "cat"]).agg(kg=("kg", "sum"), fob=("fob_tot", "sum"))
    g.index = g.index.set_names(["t", "cat"])
    KG, FOB = g.kg.unstack(), g.fob.unstack()
    P = FOB / KG
    uv = FOB.sum(axis=1) / KG.sum(axis=1)
    kb = e[e.anio == BASE].groupby("cat").kg.sum()
    w0 = kb / kb.sum()
    filas = {}
    for t in P.index:
        p = P.loc[t].dropna()
        w = w0.reindex(p.index).fillna(0.0)
        if w.sum() < .5:            # la canasta base casi no aparece: mes no usable
            continue
        w = w / w.sum()
        filas[t] = float((w * p[w.index]).sum())
    fijo = pd.Series(filas).sort_index()
    out = pd.DataFrame({"uv": uv.reindex(fijo.index), "fijo": fijo})
    out["comp"] = np.log(out.uv / out.fijo)
    part = (100 * KG.div(KG.sum(axis=1), axis=0))
    return out, P, part


# ----------------------------------------------------------------- estimación
def reestimar(s: pd.DataFrame, M: pd.DataFrame, comp_m: pd.Series) -> pd.DataFrame:
    """La misma DOLS de la elasticidad, con y sin el mix sacado del precio."""
    s = s.copy()
    s["comp"] = comp_m.rolling(12, min_periods=6).mean().reindex(
        s.index).interpolate(limit_direction="both")
    s["lp_adj"] = s.lp + s.comp          # saca del precio argentino el efecto mezcla
    K, filas = E.K_DOLS, []
    for col, nom in (("lp", "Valor unitario sin corregir"),
                     ("lp_adj", "Corregido por composición")):
        D = pd.DataFrame({f"d{k:+d}": s[col].diff().shift(k)
                          for k in range(-K, K + 1)}, index=s.index)
        X = pd.concat([s[[col, "paro"]].rename(columns={col: "lp"}), D, M],
                      axis=1).dropna()
        r = E._ols(s.lq.loc[X.index], X)
        ci = r.conf_int().loc["lp"]
        filas.append([nom, r.params["lp"], r.bse["lp"], ci[0], ci[1], int(r.nobs)])
    t = pd.DataFrame(filas, columns=["Especificación", "Elasticidad s", "EE (HAC)",
                                     "IC95 inferior", "IC95 superior", "n"])
    return t, s


# -------------------------------------------------------------------- gráfico
def grafico(part: pd.DataFrame, ia: pd.DataFrame) -> None:
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib.ticker import FuncFormatter

    coma = FuncFormatter(lambda x, _: f"{x:,.2f}".replace(".", ","))
    fig, (a1, a2) = plt.subplots(1, 2, figsize=(11.2, 4.7))
    for ax in (a1, a2):
        ax.grid(lw=.4, alpha=.4)
        for sp in ("top", "right"):
            ax.spines[sp].set_visible(False)

    orden = [c for c in ["entero L1", "entero L2", "entero L3", "cola", "pelado"]
             if c in part.columns]
    p = part[orden].fillna(0)
    a1.stackplot(p.index, *[p[c] for c in orden], labels=orden,
                 colors=[COLORES[c] for c in orden], alpha=.9)
    a1.set_ylim(0, 100)
    a1.set_ylabel("% de los kilos embarcados")
    a1.set_xlim(p.index.min(), p.index.max())
    a1.legend(frameon=False, fontsize=8, loc="lower left", ncol=3)
    nocab = [c for c in ("cola", "pelado") if c in part.columns]
    ini = float(part.loc[part.index.min(), nocab].fillna(0).sum())
    fin = float(part.loc[2025, nocab].fillna(0).sum())
    a1.set_title("La mezcla que Argentina embarca a España\n"
                 f"cola y pelado pasan de {ini:.1f}% a {fin:.1f}% de los kilos"
                 .replace(".", ","), fontsize=10.5, color=AZUL, loc="left")

    a2.plot(ia.index, ia.uv, lw=2.0, color=AZUL, marker="o", ms=4,
            label="Valor unitario observado")
    a2.plot(ia.index, ia.fijo, lw=2.0, color=ROJO, ls="--", marker="o", ms=4,
            label=f"Contrafáctico: con la mezcla de {BASE}")
    a2.fill_between(ia.index, ia.fijo, ia.uv, where=ia.uv >= ia.fijo,
                    color=AZUL, alpha=.15)
    a2.set_ylabel("US$/kg FOB")
    a2.yaxis.set_major_formatter(coma)
    a2.legend(frameon=False, fontsize=8.4, loc="lower left")
    a2.set_title("La brecha entre las dos líneas es composición\n"
                 f"+{100 * (np.exp(ia.comp.loc[2025]) - 1):.1f}% en 2025"
                 .replace(".", ","), fontsize=10.5, color=AZUL, loc="left")
    fig.text(.01, .015,
             "Fuente: elaboración propia sobre la base de comercio, despachos de "
             "exportación argentinos con destino España, calibre y presentación "
             "tomados de la descripción comercial de cada despacho. El\n"
             "contrafáctico valúa cada año con los precios de ese año pero con la "
             "participación en kilos de 2013. La cobertura de calibre va del 96% de "
             "los kilos en 2024-2026 al 36% en 2017: el índice se calcula\nsobre el "
             "subconjunto que la declara. Antes de 2020 el volumen de los despachos "
             "tiene doble conteo, que se cancela en cocientes si no discrimina por "
             "producto. 2026 es parcial: enero a julio. "
             "Elaboración: Lic. Fabián Pettigrew · AXIA.", fontsize=6.6)
    fig.tight_layout(rect=(0, .10, 1, 1))
    out = os.path.join(SAL, "espana_composicion_calibre.svg")
    fig.savefig(out)
    fig.savefig(out.replace(".svg", ".png"), dpi=200)
    plt.close(fig)
    print("\nGráfico:", out, "(+ .png)")


# ------------------------------------------------------------------- planilla
HDR = Font(name="Arial", size=10, bold=True, color="FFFFFF")
ARI = Font(name="Arial", size=10)
TIT = Font(name="Arial", size=13, bold=True, color="1F4E79")
SUB = Font(name="Arial", size=9, color="5A5A5A")
FA = PatternFill("solid", fgColor="1F4E79")
FG = PatternFill("solid", fgColor="F2F5F8")
CEN = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _hoja(wb, nombre, titulo, bajada, df, fmts, anchos):
    ws = wb.create_sheet(nombre)
    ws["A1"], ws["A1"].font = titulo, TIT
    ws["A2"], ws["A2"].font = bajada, SUB
    ws["A2"].alignment = Alignment(vertical="top", wrap_text=True)
    for j, t in enumerate(df.columns, start=1):
        c = ws.cell(4, j, str(t))
        c.font, c.fill, c.alignment = HDR, FA, CEN
    ws.row_dimensions[4].height = 34
    for i, (_, fila) in enumerate(df.iterrows()):
        for j, (v, fm) in enumerate(zip(fila, fmts), start=1):
            if isinstance(v, float) and pd.isna(v):
                v = None
            elif isinstance(v, np.integer):
                v = int(v)
            elif isinstance(v, np.floating):
                v = float(v)
            c = ws.cell(5 + i, j, v)
            c.font = ARI
            if fm:
                c.number_format = fm
                c.alignment = Alignment(horizontal="right")
            if i % 2:
                c.fill = FG
    for j, w in enumerate(anchos, start=1):
        ws.column_dimensions[L(j)].width = w
    ws.freeze_panes = "B5"


def planilla(tabs) -> None:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    F3, F2, FP = "#,##0.000", "#,##0.00", "#,##0.0"

    _hoja(wb, "Composición",
          "Efecto composición en el valor unitario argentino a España",
          "Valor unitario observado contra el contrafáctico con la mezcla congelada "
          f"en {BASE}, valuada a los precios de cada año. La diferencia es "
          "composición.",
          tabs["idx"], [None, F2, F2, F3, FP], [10, 15, 18, 15, 15])

    _hoja(wb, "Mezcla", "Participación de cada categoría en los kilos embarcados",
          "Categorías agrupadas: entero L1, L2 y L3, cola (todos los grados) y "
          "pelado. Sólo despachos con calibre declarado.",
          tabs["part"], [None] + [FP] * (len(tabs["part"].columns) - 1),
          [10] + [13] * (len(tabs["part"].columns) - 1))

    _hoja(wb, "Precios por categoría", "Precio FOB de cada categoría, US$/kg",
          "Es el precio dentro de cada casillero de producto: su movimiento es "
          "precio puro, no mezcla.",
          tabs["precios"], [None] + [F2] * (len(tabs["precios"].columns) - 1),
          [10] + [13] * (len(tabs["precios"].columns) - 1))

    _hoja(wb, "Elasticidad corregida",
          "La elasticidad de sustitución, antes y después de sacar la mezcla",
          "Misma especificación DOLS(±3) del módulo de elasticidad, cambiando sólo "
          "el precio relativo: al argentino se le saca el factor de composición.",
          tabs["est"], [None, F3, F3, F3, F3, "#,##0"], [34, 14, 12, 14, 14, 8])

    _hoja(wb, "Cobertura", "Cuánto del dato declara calibre",
          "El índice se calcula sobre el subconjunto que declara calibre en la "
          "descripción comercial. Cuanto menor la cobertura, menos representativo.",
          tabs["cob"], [None, "#,##0", FP, FP, FP], [10, 12, 15, 18, 20])

    notas = [
        ("Qué se midió", True),
        ("Si la mezcla de producto que Argentina embarca a España se corrió hacia "
         "producto más caro, el valor unitario sube sin que suba ningún precio, y la "
         "elasticidad de sustitución estimada con valores unitarios queda sesgada. "
         "Acá se separa una cosa de la otra usando el calibre y la presentación que "
         "trae la base de comercio despacho por despacho, que es lo que Comext no "
         "tiene.", False),
        ("", False),
        ("El resultado", True),
        ("La mezcla SÍ se movió: cola y pelado pasan de 0,3% de los kilos en 2013 a "
         "16,0% en 2025, y son entre 1 y 3 dólares por kilo más caros que el "
         "entero. Pero el efecto sobre el valor unitario es chico: +4,0% acumulado "
         "entre 2013 y 2025, o sea 0,039 en logaritmos, contra una divergencia del "
         "precio relativo de 0,290. La composición explica el 13,5% de esa "
         "divergencia; el 86,5% restante es precio genuino.", False),
        ("Corregida la mezcla, la elasticidad pasa de 2,17 a 2,52. Se mueve, pero "
         "queda cómodamente dentro del intervalo de confianza de la estimación "
         "original, y en la dirección esperada: si parte de la divergencia medida no "
         "era precio, la respuesta de cantidades al precio verdadero es MAYOR, no "
         "menor.", False),
        ("La conclusión económica no se mueve. Con la elasticidad corregida aplicada "
         "al cambio de precio corregido, el precio sigue explicando alrededor del "
         "45% de la pérdida de volumen relativo entre 2013 y 2025. Los dos efectos "
         "se compensan: elasticidad más alta por cambio de precio más chico.", False),
        ("", False),
        ("Por qué el efecto es chico si la mezcla cambió tanto", True),
        ("Porque el entero sigue siendo el grueso: L1 más L2 más L3 son el 84% de "
         "los kilos en 2025. Cola y pelado crecieron mucho en términos relativos "
         "pero desde casi cero, y 16% de los kilos a un sobreprecio de 1 a 3 dólares "
         "mueve el promedio poco. Además el efecto no es una deriva sostenida: sube "
         "a 10,3% en 2023 y vuelve a 4,0% en 2025. Es volatilidad de mezcla, no un "
         "corrimiento estructural.", False),
        ("", False),
        ("Límites", True),
        ("La cobertura de calibre va del 96% de los kilos en 2024-2026 al 36% en "
         "2017. El índice se calcula sobre lo que declara y se supone que el resto "
         "se comporta igual.", False),
        ("La base de comercio es FOB en Argentina y Comext es CIF en frontera "
         "española: el factor se mide en una punta y se aplica en la otra. La "
         "correlación entre las variaciones anuales de los dos valores unitarios es "
         "0,39. La corrección es aproximada.", False),
        ("Antes de 2020 el volumen de los despachos tiene doble conteo. Las "
         "participaciones son cocientes, así que se cancela si la duplicación no "
         "discrimina por tipo de producto; no hay forma de verificarlo.", False),
        ("Dentro de «cola» el grado no se declara de manera consistente, así que no "
         "se puede abrir más fino sin perder comparabilidad.", False),
        ("", False),
        ("Reproducibilidad", True),
        ("«composicion_calibre_es.py», que usa «fuentes/expo.py» para la base de "
         "comercio y «elasticidad_sustitucion_es.py» para la especificación. "
         "Elaboración: Lic. Fabián Pettigrew · AXIA.", False),
    ]
    ws = wb.create_sheet("Notas")
    ws.column_dimensions["A"].width = 118
    for i, (t, neg) in enumerate(notas, start=1):
        c = ws.cell(i, 1, t)
        c.font = Font(name="Arial", size=10, bold=neg,
                      color="1F4E79" if neg else "1A1A1A")
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if not neg and t:
            ws.row_dimensions[i].height = 14 * (1 + len(t) // 108)

    for h in wb.sheetnames:
        wb[h].sheet_properties.tabColor = "7F7F7F" if h in ("Notas", "Cobertura") \
            else "1F4E79"
    out = os.path.join(SAL, "Composicion_calibre_Espana.xlsx")
    wb.save(out)
    print("Planilla:", out)


# ----------------------------------------------------------------------- main
def main() -> None:
    print("=" * 98)
    print("COMPOSICIÓN POR CALIBRE: ¿CUÁNTO DE LA DIVERGENCIA DE PRECIO ES MEZCLA "
          "DE PRODUCTO?")
    print("=" * 98)
    e = despachos()
    ia, Pa, part = indice(e, "A")
    im, _, _ = indice(e, "M")

    print("\n" + "-" * 98 + "\nMEZCLA EMBARCADA A ESPAÑA (% de los kilos)\n" + "-" * 98)
    print(part.round(1).to_string())

    print("\n" + "-" * 98 + "\nVALOR UNITARIO OBSERVADO CONTRA LA CANASTA "
          f"CONGELADA EN {BASE}\n" + "-" * 98)
    t = ia.copy()
    t["% composición"] = 100 * (np.exp(t.comp) - 1)
    print(t.round(3).to_string())

    c25 = float(ia.comp.loc[2025])
    print("\nEfecto composición acumulado 2013 → 2025: "
          f"{c25:+.3f} en logaritmos ({100 * (np.exp(c25) - 1):+.1f}%)")
    print(f"Divergencia del precio relativo en Comext 2013 → 2025: -0,290")
    print(f"  ⇒ la composición explica {100 * c25 / 0.290:.1f}% de la divergencia; "
          f"el resto, {100 - 100 * c25 / 0.290:.1f}%, es precio genuino")

    s, M = E.panel()
    est, s2 = reestimar(s, M, im.comp)
    print("\n" + "-" * 98 + "\nLA ELASTICIDAD, ANTES Y DESPUÉS DE SACAR LA MEZCLA\n"
          + "-" * 98)
    print(est.round(3).to_string(index=False))

    # la conclusión económica: ¿cambia el reparto entre precio y no-precio?
    a = s2.assign(v_ar=s2.q_ar * s2.p_ar, v_ec=s2.q_ec * s2.p_ec)
    a = a.groupby(a.index.year).agg(q_ar=("q_ar", "sum"), q_ec=("q_ec", "sum"),
                                    v_ar=("v_ar", "sum"), v_ec=("v_ec", "sum"))
    dlq = float(np.log(a.q_ar[2025] / a.q_ec[2025])
                - np.log(a.q_ar[2013] / a.q_ec[2013]))
    dlp = -0.290
    dlp_adj = dlp + c25
    s0, s1 = est["Elasticidad s"].iloc[0], est["Elasticidad s"].iloc[1]
    print("\n" + "-" * 98 + "\n¿CAMBIA LA CONCLUSIÓN?\n" + "-" * 98)
    print(f"Sin corregir:  s = {s0:.2f} x cambio de precio {dlp:+.3f} = "
          f"{s0 * dlp:+.3f}  ⇒ explica {100 * s0 * dlp / dlq:.1f}% de la pérdida "
          f"de volumen relativo")
    print(f"Corregido:     s = {s1:.2f} x cambio de precio {dlp_adj:+.3f} = "
          f"{s1 * dlp_adj:+.3f}  ⇒ explica {100 * s1 * dlp_adj / dlq:.1f}%")
    print("Los dos efectos se compensan: elasticidad más alta sobre un cambio de "
          "precio más chico.")

    tabs = {
        "idx": ia.reset_index().rename(columns={
            "index": "Año", "t": "Año", "uv": "Valor unitario observado",
            "fijo": f"Con la mezcla de {BASE}", "comp": "ln(composición)"}).assign(
                **{"% composición": 100 * (np.exp(ia.comp.values) - 1)}),
        "part": part.round(2).reset_index().rename(columns={"t": "Año"}),
        "precios": Pa.round(3).reset_index().rename(columns={"t": "Año"}),
        "est": est,
        "cob": cobertura(e),
    }
    grafico(part, ia)
    planilla(tabs)


if __name__ == "__main__":
    main()
