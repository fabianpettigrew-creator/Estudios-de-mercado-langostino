# -*- coding: utf-8 -*-
"""Elasticidad de sustitución entre el langostino argentino y el vannamei ecuatoriano
en el mercado español de camarón congelado, 2013-01 a 2026-07.

LA PREGUNTA. La cuota argentina en España cayó de 29,4% en 2024 a 20,9% en 2025
mientras la brecha de precio contra el vannamei se abría de +1,33 a +1,61 EUR/kg.
¿Cuánto de esa pérdida es precio y cuánto es otra cosa?

EL MODELO. Demanda relativa tipo Armington/CES entre dos variedades del mismo bien:

    ln(q_ar / q_ec) = a + s · ln(p_ec / p_ar) + estacionalidad + u

donde `s` es la elasticidad de sustitución. Si s = 2, abrir un 10% la prima de precio
argentina cuesta 20% del volumen relativo. El signo esperado es positivo: cuanto más
caro está el langostino contra el vannamei, menos langostino se compra.

POR QUÉ NO ALCANZA CON CORRER UN MÍNIMO CUADRADO Y LISTO. Las dos series son I(1) y
están tendenciadas —la correlación de la variable de precio con una tendencia lineal es
-0,84—, así que una regresión en niveles puede ser espuria. Metiendo tendencia el
coeficiente se derrumba a 0,4 y deja de ser significativo, pero eso no prueba que no
haya sustitución: la tendencia y el precio relativo son casi la misma variable y meter
las dos es pedirle al dato que distinga lo que no puede. La salida correcta es testear
cointegración. Si las series cointegran, la relación de largo plazo existe, la
tendencia determinística sobra, y el estimador indicado es DOLS, que corrige de una vez
la endogeneidad del precio y la autocorrelación del residuo.

POR QUÉ NO HAY VARIABLES INSTRUMENTALES. Se probaron dos instrumentos de oferta: el
paro de 2025 y el precio FOB argentino a destinos distintos de España (instrumento tipo
Hausman). El primero da s = 7,3, que es la razón de Wald entre el derrumbe de volumen y
la suba de precio de un único episodio que coincidió con la expansión ecuatoriana: le
atribuye al precio todo lo que pasó en 2025. El segundo es débil (F de primera etapa =
4,1) y su intervalo va de 0,3 a 6,4. Corridos juntos, Sargan rechaza con p = 0,002: los
dos instrumentos no pueden ser válidos a la vez. Están reportados en la planilla para
que se vea el ejercicio, pero no se usan para concluir. En una regresión cointegrante
DOLS ya resuelve la endogeneidad, que es para lo que se lo usa acá.

QUÉ NO MIDE ESTO. Las dos puntas no son el mismo producto: langostino salvaje entero y
cola contra vannamei de cultivo. El nivel del precio relativo mezcla precio con
producto y calibre. Lo que el modelo lee es la RESPUESTA DE CANTIDADES A CAMBIOS en ese
precio relativo, no una comparación de niveles, y ahí la mezcla de producto se va en la
constante y en las dummies mensuales mientras no cambie sistemáticamente. Si la mezcla
argentina se corrió hacia producto más caro a lo largo de la serie —cosa que el dato de
calibre no permite descartar—, parte de lo que el modelo lee como precio es composición.

Fuente: Eurostat-Comext DS-045409, declarante España, NC 0306.16 y 0306.17 congelado,
mensual. Precio = valor unitario CIF en frontera española.

Uso:  python elasticidad_sustitucion_es.py
Salida: consola
        salidas/Elasticidad_sustitucion_AR_EC_Espana.xlsx
        salidas/espana_elasticidad_sustitucion.svg  (+ .png)
"""
from __future__ import annotations

import os
import sys
import warnings

import numpy as np
import openpyxl
import pandas as pd
import statsmodels.api as sm
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter as L
from statsmodels.tsa.stattools import adfuller, coint, grangercausalitytests, kpss
from statsmodels.tsa.vector_ar.vecm import coint_johansen

warnings.filterwarnings("ignore")
sys.stdout.reconfigure(encoding="utf-8")
RAIZ = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, RAIZ)
SAL = os.path.join(RAIZ, "salidas")

import comext_espana_cuota_brecha as C   # noqa: E402  (carga la base de Comext)

HAC = 6                                              # Newey-West, serie mensual
PARO = (pd.Period("2025-04"), pd.Period("2025-11"))  # conflicto gremial y rebote
K_DOLS = 3                                           # adelantos y rezagos de Stock-Watson
AZUL, ROJO, GRIS = "#1f4e79", "#c0504d", "#9aa7ad"


# ---------------------------------------------------------------------- datos
def panel() -> tuple[pd.DataFrame, pd.DataFrame]:
    """Serie mensual de cantidades y precios relativos, y dummies de mes."""
    m = C.mensual(C.cargar())
    s = pd.DataFrame(index=m.index)
    s["q_ar"], s["q_ec"] = m.ar_t, m.ec_t
    s["p_ar"], s["p_ec"] = m.ar_p, m.ec_p
    s["lq"] = np.log(m.ar_t) - np.log(m.ec_t)      # volumen relativo
    s["lp"] = np.log(m.ec_p) - np.log(m.ar_p)      # precio relativo: s > 0 esperado
    s["paro"] = ((s.index >= PARO[0]) & (s.index <= PARO[1])).astype(float)
    if (s[["q_ar", "q_ec"]] <= 0).any().any():
        sys.exit("ABORTA: hay meses con volumen cero; el logaritmo no existe")
    M = pd.get_dummies(s.index.month, prefix="m",
                       drop_first=True).set_index(s.index).astype(float)
    # lq sin estacionalidad, sólo para los tests de raíz unitaria y cointegración
    s["lq_sa"] = (sm.OLS(s.lq, sm.add_constant(M)).fit().resid + s.lq.mean())
    return s, M


def instrumento_hausman(idx) -> pd.Series:
    """Precio FOB del langostino argentino a destinos distintos de España.

    Desplaza el costo de la oferta argentina sin pasar por la demanda española. Sale
    de la aduana argentina, que arranca en 2017-02 en esta base."""
    d = pd.read_pickle(os.path.join(RAIZ, "datos", "aduana_langostino.pkl"))
    d["pais_cod"] = d.pais_cod.astype(str)
    o = d[d.pais_cod != "410"].copy()          # 410 = España
    o["per"] = pd.PeriodIndex(o.anio.astype(str) + "-"
                              + o.mes.astype(str).str.zfill(2), freq="M")
    g = o.groupby("per")[["kg", "fob"]].sum()
    return np.log(g.fob / g.kg).reindex(idx)


# -------------------------------------------------------------- raíces y coint
def raices(s: pd.DataFrame) -> pd.DataFrame:
    filas = []
    for c, nom in (("lq_sa", "ln(q_ar/q_ec) desestacionalizado"),
                   ("lp", "ln(p_ec/p_ar)")):
        x = s[c].dropna()
        a, k = adfuller(x, regression="c", autolag="AIC"), kpss(x, regression="c",
                                                                nlags="auto")
        da = adfuller(x.diff().dropna(), regression="c", autolag="AIC")
        dk = kpss(x.diff().dropna(), regression="c", nlags="auto")
        filas.append([nom, a[0], a[1], k[0], k[1], da[0], da[1], dk[1],
                      "I(1)" if (a[1] > .05 and da[1] < .05) else "revisar"])
    return pd.DataFrame(filas, columns=["Serie", "ADF nivel", "ADF p", "KPSS nivel",
                                        "KPSS p", "ADF dif", "ADF dif p",
                                        "KPSS dif p", "Conclusión"])


def cointegracion(s: pd.DataFrame) -> tuple[pd.DataFrame, float]:
    eg1 = coint(s.lq_sa, s.lp, trend="c", autolag="AIC")
    eg2 = coint(s.lp, s.lq_sa, trend="c", autolag="AIC")
    j = coint_johansen(s[["lq_sa", "lp"]].values, det_order=0, k_ar_diff=2)
    beta = float(-j.evec[1, 0] / j.evec[0, 0])
    gr = grangercausalitytests(s[["lq_sa", "lp"]].diff().dropna(), maxlag=3,
                               verbose=False)
    gr2 = grangercausalitytests(s[["lp", "lq_sa"]].diff().dropna(), maxlag=3,
                                verbose=False)
    filas = [
        ["Engle-Granger, lq como dependiente", eg1[0], eg1[1],
         "cointegran al 5%" if eg1[1] < .05 else "no rechaza"],
        ["Engle-Granger, lp como dependiente", eg2[0], eg2[1],
         "cointegran al 5%" if eg2[1] < .05 else "no rechaza"],
        ["Johansen, traza r=0", j.lr1[0], np.nan,
         f"rechaza al 5% (crítico {j.cvt[0, 1]:.2f})" if j.lr1[0] > j.cvt[0, 1]
         else "no rechaza"],
        ["Johansen, traza r<=1", j.lr1[1], np.nan,
         f"no rechaza (crítico {j.cvt[1, 1]:.2f})" if j.lr1[1] < j.cvt[1, 1]
         else "rechaza"],
        ["Johansen, beta normalizado sobre lq", beta, np.nan,
         "es la elasticidad de largo plazo que implica Johansen"],
        ["Granger: el precio relativo causa al volumen relativo", np.nan,
         min(gr[i + 1][0]["ssr_ftest"][1] for i in range(3)),
         "p mínimo entre 1 y 3 rezagos"],
        ["Granger: el volumen relativo causa al precio relativo", np.nan,
         min(gr2[i + 1][0]["ssr_ftest"][1] for i in range(3)),
         "p mínimo entre 1 y 3 rezagos"],
    ]
    return pd.DataFrame(filas, columns=["Prueba", "Estadístico", "p", "Lectura"]), beta


# ------------------------------------------------------------------ estimación
def _ols(y, X):
    return sm.OLS(y, sm.add_constant(X)).fit(cov_type="HAC",
                                             cov_kwds={"maxlags": HAC})


def _fila(nom, r, n=None, extra=""):
    ci = r.conf_int()
    lo, hi = (ci.loc["lp", 0], ci.loc["lp", 1]) if hasattr(ci, "loc") else (np.nan,) * 2
    return [nom, r.params["lp"], r.bse["lp"], r.tvalues["lp"], lo, hi,
            int(n if n is not None else r.nobs), extra]


def estimaciones(s: pd.DataFrame, M: pd.DataFrame) -> tuple[pd.DataFrame, dict]:
    filas, guarda = [], {}

    # 1. OLS estático, con y sin control del paro, y con tendencia para mostrar
    #    por qué la tendencia no va
    r = _ols(s.lq, pd.concat([s[["lp"]], M], axis=1))
    filas.append(_fila("OLS estático", r, extra="sin controles de oferta"))
    r = _ols(s.lq, pd.concat([s[["lp", "paro"]], M], axis=1))
    filas.append(_fila("OLS + dummy del paro", r, extra="control del shock de 2025"))
    guarda["ols"] = r
    rt = _ols(s.lq, pd.concat([s[["lp", "paro"]].assign(trend=np.arange(len(s))),
                               M], axis=1))
    filas.append(_fila("OLS + tendencia lineal", rt,
                       extra="colinealidad: corr(lp, trend) = "
                             f"{np.corrcoef(s.lp, np.arange(len(s)))[0, 1]:.2f}"))

    # 2. DOLS de Stock-Watson: adelantos y rezagos de la diferencia del precio
    D = pd.DataFrame({f"dlp{k:+d}": s.lp.diff().shift(k)
                      for k in range(-K_DOLS, K_DOLS + 1)}, index=s.index)
    X = pd.concat([s[["lp", "paro"]], D, M], axis=1).dropna()
    rd = _ols(s.lq.loc[X.index], X)
    filas.append(_fila(f"DOLS(±{K_DOLS})  ←  estimación principal", rd,
                       extra="corrige endogeneidad y autocorrelación"))
    guarda["dols"] = rd

    # 2b. La regresión inversa. Granger dice que la causalidad corre de cantidad a
    #     precio, que es la estructura de demanda inversa que usa el resto del
    #     estudio. Estimada en esa dirección, 1/b es la cota superior de s.
    Di = pd.DataFrame({f"dlq{k:+d}": s.lq.diff().shift(k)
                       for k in range(-K_DOLS, K_DOLS + 1)}, index=s.index)
    Xi = pd.concat([s[["lq", "paro"]], Di, M], axis=1).dropna()
    ri = _ols(s.lp.loc[Xi.index], Xi)
    b, eb = ri.params["lq"], ri.bse["lq"]
    ci = ri.conf_int().loc["lq"]
    filas.append([f"DOLS inverso: lp sobre lq  →  s = 1/b", 1 / b, eb / b ** 2,
                  ri.tvalues["lq"], 1 / ci[1], 1 / ci[0], int(ri.nobs),
                  f"b = {b:.3f}; cota superior, atenuada por error de medición"])
    guarda["inv"] = ri

    # 3. Robustez
    sub = s[s.index < pd.Period("2025-01")]
    Ms = M.loc[sub.index]
    r = _ols(sub.lq, pd.concat([sub[["lp"]], Ms], axis=1))
    filas.append(_fila("Robustez: 2013-2024, sin el shock", r,
                       extra="descarta que todo venga de 2025"))
    sub2 = s[s.index >= pd.Period("2018-01")]
    r = _ols(sub2.lq, pd.concat([sub2[["lp", "paro"]], M.loc[sub2.index]], axis=1))
    filas.append(_fila("Robustez: 2018-2026", r, extra="desde el cambio de signo"))

    q = s.groupby([s.index.year, (s.index.month - 1) // 3 + 1]).agg(
        q_ar=("q_ar", "sum"), q_ec=("q_ec", "sum"),
        v_ar=("q_ar", "sum"), v_ec=("q_ec", "sum"))
    qa = s.assign(v_ar=s.q_ar * s.p_ar * 1000, v_ec=s.q_ec * s.p_ec * 1000)
    q = qa.groupby([qa.index.year, (qa.index.month - 1) // 3 + 1]).agg(
        q_ar=("q_ar", "sum"), q_ec=("q_ec", "sum"),
        v_ar=("v_ar", "sum"), v_ec=("v_ec", "sum"))
    q["lq"] = np.log(q.q_ar) - np.log(q.q_ec)
    q["lp"] = np.log(q.v_ec / q.q_ec) - np.log(q.v_ar / q.q_ar)
    Q = pd.get_dummies([i[1] for i in q.index], prefix="t",
                       drop_first=True).set_index(q.index).astype(float)
    r = _ols(q.lq, pd.concat([q[["lp"]], Q], axis=1))
    filas.append(_fila("Robustez: datos trimestrales", r,
                       extra="menos ruido de valor unitario"))
    guarda["trim"] = q

    tabla = pd.DataFrame(filas, columns=["Especificación", "Elasticidad s", "EE (HAC)",
                                         "t", "IC95 inferior", "IC95 superior",
                                         "n", "Comentario"])
    return tabla, guarda


def iv(s: pd.DataFrame, M: pd.DataFrame) -> pd.DataFrame:
    """El ejercicio instrumental que NO se usa para concluir. Queda documentado."""
    from linearmodels.iv import IV2SLS
    b = pd.concat([s, M], axis=1)
    b["pterc"] = instrumento_hausman(s.index)
    b["const"] = 1.0
    b.index = b.index.to_timestamp()
    ex = ["const"] + M.columns.tolist()
    filas = []
    for instr, nom in ((["paro"], "IV · dummy del paro 2025"),
                       (["pterc"], "IV · precio FOB a terceros destinos"),
                       (["paro", "pterc"], "IV · los dos juntos")):
        dd = b[["lq", "lp"] + ex + instr].dropna()
        r = IV2SLS(dd.lq, dd[ex], dd[["lp"]], dd[instr]).fit(
            cov_type="kernel", kernel="bartlett", bandwidth=HAC)
        ci = r.conf_int().loc["lp"]
        f = float(r.first_stage.diagnostics["f.stat"].iloc[0])
        sar = float(r.sargan.pval) if len(instr) > 1 else np.nan
        filas.append([nom, float(r.params.lp), float(r.std_errors.lp), ci.lower,
                      ci.upper, f, sar, int(r.nobs)])
    return pd.DataFrame(filas, columns=["Especificación", "Elasticidad s", "EE",
                                        "IC95 inferior", "IC95 superior",
                                        "F 1ra etapa", "Sargan p", "n"])


def ecm(s: pd.DataFrame, M: pd.DataFrame, r_lp) -> tuple[pd.DataFrame, pd.Series]:
    """Corrección de error: cuánto de la sustitución ocurre mes a mes."""
    resid = s.lq - r_lp.predict(sm.add_constant(
        pd.concat([s[["lp", "paro"]], M], axis=1)))
    E = pd.DataFrame(index=s.index)
    E["dlq"], E["dlp"] = s.lq.diff(), s.lp.diff()
    E["dlp1"], E["dlq1"] = s.lp.diff().shift(1), s.lq.diff().shift(1)
    E["ec1"], E["paro"] = resid.shift(1), s.paro
    X = pd.concat([E[["dlp", "dlp1", "dlq1", "ec1", "paro"]], M], axis=1).dropna()
    r = _ols(E.dlq.loc[X.index], X)
    vm = np.log(.5) / np.log(1 + r.params["ec1"])
    filas = [
        ["Elasticidad de corto plazo (impacto del mes)", r.params["dlp"],
         r.bse["dlp"], r.tvalues["dlp"]],
        ["Velocidad de ajuste al equilibrio (alfa)", r.params["ec1"],
         r.bse["ec1"], r.tvalues["ec1"]],
        ["Vida media de una desviación, en meses", vm, np.nan, np.nan],
    ]
    return (pd.DataFrame(filas, columns=["Concepto", "Valor", "EE (HAC)", "t"]),
            resid)


def descomposicion(s: pd.DataFrame, sigma: float) -> pd.DataFrame:
    """De la caída del volumen relativo, ¿cuánto explica el precio?"""
    a = s.groupby(s.index.year).agg(q_ar=("q_ar", "sum"), q_ec=("q_ec", "sum"),
                                    v_ar=("q_ar", "sum"), v_ec=("q_ec", "sum"))
    va = s.assign(v_ar=s.q_ar * s.p_ar, v_ec=s.q_ec * s.p_ec)
    a = va.groupby(va.index.year).agg(q_ar=("q_ar", "sum"), q_ec=("q_ec", "sum"),
                                      v_ar=("v_ar", "sum"), v_ec=("v_ec", "sum"))
    a["lq"] = np.log(a.q_ar) - np.log(a.q_ec)
    a["lp"] = np.log(a.v_ec / a.q_ec) - np.log(a.v_ar / a.q_ar)
    filas = []
    for ini, fin, nom in ((2013, 2025, "2013 → 2025, toda la serie"),
                          (2015, 2025, "2015 → 2025, desde el pico argentino"),
                          (2024, 2025, "2024 → 2025, el año del derrumbe")):
        dlq, dlp = a.lq[fin] - a.lq[ini], a.lp[fin] - a.lp[ini]
        pred = sigma * dlp
        filas.append([nom, dlp, dlq, pred, dlq - pred,
                      100 * pred / dlq if dlq else np.nan])
    return pd.DataFrame(filas, columns=["Tramo", "Cambio en ln(p_ec/p_ar)",
                                        "Cambio observado en ln(q_ar/q_ec)",
                                        "Explicado por precio", "Residuo",
                                        "% explicado por precio"])


# -------------------------------------------------------------------- gráfico
def grafico(s: pd.DataFrame, sigma: float, const: float) -> None:
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib.ticker import FuncFormatter

    coma = FuncFormatter(lambda x, _: f"{x:,.2f}".replace(".", ","))
    a = s.assign(v_ar=s.q_ar * s.p_ar, v_ec=s.q_ec * s.p_ec)
    a = a.groupby(a.index.year).agg(q_ar=("q_ar", "sum"), q_ec=("q_ec", "sum"),
                                    v_ar=("v_ar", "sum"), v_ec=("v_ec", "sum"))
    a["lq"] = np.log(a.q_ar) - np.log(a.q_ec)
    a["lp"] = np.log(a.v_ec / a.q_ec) - np.log(a.v_ar / a.q_ar)

    fig, (a1, a2) = plt.subplots(1, 2, figsize=(11.2, 4.9),
                                 gridspec_kw={"width_ratios": [1.25, 1]})
    for ax in (a1, a2):
        ax.grid(lw=.4, alpha=.4)
        for sp in ("top", "right"):
            ax.spines[sp].set_visible(False)

    x = s.index.to_timestamp()
    r12 = s[["lq", "lp"]].rolling(12).mean()
    a1.plot(x, r12.lq, lw=2.0, color=AZUL, label="Volumen relativo  ln(q_ar/q_ec)")
    a1.axhline(0, lw=.8, color="#7f7f7f")
    a1.set_ylabel("ln(q_ar / q_ec)", color=AZUL)
    a1.yaxis.set_major_formatter(coma)
    b = a1.twinx()
    b.plot(x, r12.lp, lw=2.0, color=ROJO, label="Precio relativo  ln(p_ec/p_ar)")
    b.set_ylabel("ln(p_ec / p_ar)", color=ROJO)
    b.yaxis.set_major_formatter(coma)
    b.spines["top"].set_visible(False)
    h1, l1 = a1.get_legend_handles_labels()
    h2, l2 = b.get_legend_handles_labels()
    a1.legend(h1 + h2, l1 + l2, frameon=False, fontsize=8.4, loc="lower left")
    a1.set_title("Las dos series se mueven juntas\nmedias móviles de 12 meses",
                 fontsize=10.5, color=AZUL, loc="left")

    sc = a2.scatter(a.lp, a.lq, c=a.index, cmap="viridis", s=64, zorder=3,
                    edgecolor="white", lw=.8)
    xs = np.linspace(a.lp.min() - .03, a.lp.max() + .03, 50)
    # la constante de la regresión mensual corresponde a enero sin paro, así que
    # para dibujarla contra los promedios anuales se recentra en el centroide
    b0 = a.lq.mean() - sigma * a.lp.mean()
    a2.plot(xs, b0 + sigma * xs, lw=1.6, color=ROJO, zorder=2,
            label=f"pendiente DOLS = {sigma:.2f}".replace(".", ","))
    for y0 in (2013, 2020, 2024, 2025):
        if y0 in a.index:
            a2.annotate(str(y0), (a.lp[y0], a.lq[y0]), xytext=(7, -3),
                        textcoords="offset points", fontsize=8, color="#404040")
    a2.set_xlabel("Precio relativo  ln(p_ec / p_ar)")
    a2.set_ylabel("Volumen relativo  ln(q_ar / q_ec)")
    a2.xaxis.set_major_formatter(coma)
    a2.yaxis.set_major_formatter(coma)
    a2.legend(frameon=False, fontsize=8.6, loc="lower right")
    a2.set_title("Promedios anuales y recta de largo plazo\n"
                 "cada punto es un año, el color va de 2013 a 2026",
                 fontsize=10.5, color=AZUL, loc="left")
    fig.colorbar(sc, ax=a2, pad=.02).ax.tick_params(labelsize=7)

    fig.text(.01, .015,
             "Fuente: elaboración propia sobre Eurostat-Comext DS-045409, declarante "
             "España, NC 0306.16 y 0306.17 congelado, mensual 2013-01 a 2026-07. "
             "Elasticidad de sustitución estimada por DOLS(±3) sobre la\nrelación "
             "cointegrante, con dummies mensuales y control del paro de 2025. El "
             "precio es valor unitario CIF: las dos puntas no son el mismo producto, "
             "así que el modelo lee respuesta a CAMBIOS en el precio\nrelativo, no "
             "niveles. Elaboración: Lic. Fabián Pettigrew · AXIA.", fontsize=6.6)
    fig.tight_layout(rect=(0, .10, 1, 1))
    out = os.path.join(SAL, "espana_elasticidad_sustitucion.svg")
    fig.savefig(out)
    fig.savefig(out.replace(".svg", ".png"), dpi=200)
    plt.close(fig)
    print("\nGráfico:", out, "(+ .png)")


# ------------------------------------------------------------------- planilla
HDR = Font(name="Arial", size=10, bold=True, color="FFFFFF")
ARI = Font(name="Arial", size=10)
TIT = Font(name="Arial", size=13, bold=True, color="1F4E79")
SUB = Font(name="Arial", size=9, color="5A5A5A")
FA = PatternFill("solid", fgColor="1F4E79")
FG = PatternFill("solid", fgColor="F2F5F8")
CEN = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _hoja(wb, nombre, titulo, bajada, df, fmts, anchos):
    ws = wb.create_sheet(nombre)
    ws["A1"], ws["A1"].font = titulo, TIT
    ws["A2"], ws["A2"].font = bajada, SUB
    ws["A2"].alignment = Alignment(vertical="top", wrap_text=True)
    for j, t in enumerate(df.columns, start=1):
        c = ws.cell(4, j, str(t))
        c.font, c.fill, c.alignment = HDR, FA, CEN
    ws.row_dimensions[4].height = 34
    for i, (_, fila) in enumerate(df.iterrows()):
        for j, (v, fm) in enumerate(zip(fila, fmts), start=1):
            if isinstance(v, (np.floating, float)) and pd.isna(v):
                v = None
            elif isinstance(v, (np.integer,)):
                v = int(v)
            elif isinstance(v, (np.floating,)):
                v = float(v)
            c = ws.cell(5 + i, j, v)
            c.font = ARI
            if fm:
                c.number_format = fm
                c.alignment = Alignment(horizontal="right")
            if i % 2:
                c.fill = FG
    for j, w in enumerate(anchos, start=1):
        ws.column_dimensions[L(j)].width = w
    ws.freeze_panes = "B5"


def planilla(s, tabs) -> None:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    F3, F2, FP = "#,##0.000", "#,##0.00", "#,##0.0"

    _hoja(wb, "Estimaciones",
          "Elasticidad de sustitución langostino argentino / vannamei ecuatoriano "
          "en España",
          "Demanda relativa ln(q_ar/q_ec) = a + s·ln(p_ec/p_ar) + dummies de mes. "
          "Errores estándar HAC de Newey-West con 6 rezagos. La fila marcada es la "
          "estimación que se reporta.",
          tabs["est"], [None, F3, F3, F2, F3, F3, "#,##0", None],
          [40, 13, 12, 9, 13, 13, 7, 44])

    _hoja(wb, "Raíces unitarias", "¿Son estacionarias las series?",
          "ADF contrasta H0 = raíz unitaria (p alto ⇒ no estacionaria). KPSS "
          "contrasta H0 = estacionariedad (p bajo ⇒ no estacionaria). Las dos "
          "pruebas coinciden: las series son I(1) en nivel y estacionarias en "
          "primera diferencia, que es la condición para buscar cointegración.",
          tabs["raices"], [None, F3, F3, F3, F3, F3, F3, F3, None],
          [36, 11, 10, 11, 10, 11, 11, 11, 12])

    _hoja(wb, "Cointegración", "¿Existe la relación de largo plazo?",
          "Si las series cointegran, la regresión en niveles no es espuria y la "
          "tendencia determinística sobra. Engle-Granger rechaza con lq como "
          "dependiente; Johansen encuentra exactamente un vector cointegrante.",
          tabs["coint"], [None, F3, F3, None], [46, 13, 10, 48])

    _hoja(wb, "Corto plazo (MCE)", "Modelo de corrección de error",
          "Separa lo que pasa mes a mes de la relación de largo plazo. La "
          "elasticidad de impacto no es significativa: la sustitución no ocurre "
          "dentro del mes.",
          tabs["ecm"], [None, F3, F3, F2], [46, 12, 12, 9])

    _hoja(wb, "Descomposición", "De la caída del volumen relativo, ¿cuánto es precio?",
          "Aplica la elasticidad estimada al cambio observado del precio relativo y "
          "compara contra el cambio observado del volumen relativo. El residuo es lo "
          "que el precio no explica: disponibilidad, formato, empuje comercial, "
          "contratos.",
          tabs["desc"], [None, F3, F3, F3, F3, FP], [34, 17, 18, 15, 12, 15])

    _hoja(wb, "Variables instrumentales",
          "El ejercicio de VI, que NO se usa para concluir",
          "Queda documentado para que se vea por qué se descarta: el paro da una "
          "elasticidad de 7,3 atribuyéndole al precio todo el derrumbe de 2025; el "
          "instrumento de costo es débil (F = 4,1); corridos juntos, Sargan rechaza. "
          "En una regresión cointegrante DOLS ya corrige la endogeneidad.",
          tabs["iv"], [None, F3, F3, F3, F3, F2, F3, "#,##0"],
          [34, 13, 10, 13, 13, 12, 11, 7])

    ser = s.reset_index()
    ser.columns = ["Mes"] + list(ser.columns[1:])
    ser["Mes"] = [f"{p.year}-{p.month:02d}" for p in s.index]
    _hoja(wb, "Serie", "Serie mensual usada en la estimación",
          "q en toneladas, p en EUR/kg CIF. lq y lp son los logaritmos de los "
          "cocientes. lq_sa es lq sin estacionalidad, que se usa sólo en los tests.",
          ser[["Mes", "q_ar", "q_ec", "p_ar", "p_ec", "lq", "lp", "paro", "lq_sa"]],
          [None, FP, FP, F2, F2, F3, F3, "0", F3],
          [10, 11, 11, 10, 10, 10, 10, 7, 10])

    notas = [
        ("Qué se estima", True),
        ("La elasticidad de sustitución entre dos variedades del mismo bien en un "
         "mercado: el langostino argentino y el vannamei ecuatoriano puestos en "
         "España. El modelo es la demanda relativa de Armington/CES, "
         "ln(q_ar/q_ec) = a + s·ln(p_ec/p_ar), donde s responde a: si el langostino "
         "se encarece 10% contra el vannamei, ¿cuánto volumen relativo pierde?", False),
        ("", False),
        ("Por qué no se mete tendencia", True),
        ("Las dos series son I(1) y están cointegradas: Engle-Granger rechaza la no "
         "cointegración al 5% y Johansen encuentra exactamente un vector. Con "
         "cointegración la regresión en niveles no es espuria y el estimador es "
         "superconsistente. Meter además una tendencia lineal derrumba el "
         "coeficiente a 0,4 y lo vuelve no significativo, pero eso no es evidencia "
         "de que no haya sustitución: la correlación entre el precio relativo y la "
         "tendencia es -0,84, o sea que son casi la misma variable. La "
         "especificación con tendencia está en la planilla para que se vea, no para "
         "que se cite.", False),
        ("", False),
        ("Por qué no se usan variables instrumentales", True),
        ("El precio relativo es endógeno: oferta y demanda lo determinan juntas. La "
         "respuesta habitual son instrumentos de oferta, y se probaron dos. El paro "
         "de 2025 da s = 7,3, que es la razón de Wald de un episodio único que "
         "coincidió con la expansión ecuatoriana: le carga al precio todo lo que "
         "pasó ese año. El precio FOB argentino a terceros destinos es débil, con F "
         "de primera etapa de 4,1. Corridos juntos, Sargan rechaza con p = 0,002, "
         "que es la prueba de que no pueden ser los dos válidos. En una regresión "
         "cointegrante, DOLS de Stock-Watson ya corrige la endogeneidad y la "
         "autocorrelación, y es lo que se usa.", False),
        ("", False),
        ("Qué NO mide", True),
        ("Las dos puntas no son el mismo producto: salvaje entero y cola contra "
         "vannamei de cultivo. El NIVEL del precio relativo mezcla precio con "
         "producto y calibre. El modelo lee respuesta a CAMBIOS, y la mezcla se va "
         "en la constante mientras no se corra sistemáticamente. Si la oferta "
         "argentina se corrió hacia producto más caro a lo largo de la serie, parte "
         "de lo que el modelo atribuye al precio es composición. El dato de calibre "
         "de Comext no permite descartarlo.", False),
        ("El valor unitario mensual tiene error de medición, que atenúa el "
         "coeficiente hacia cero. La estimación trimestral, con menos ruido, está "
         "entre las de robustez.", False),
        ("Los datos de 2025 y 2026 son provisionales.", False),
        ("", False),
        ("Reproducibilidad", True),
        ("«elasticidad_sustitucion_es.py», que se apoya en "
         "«comext_espana_cuota_brecha.py» para cargar la base. "
         "Elaboración: Lic. Fabián Pettigrew · AXIA.", False),
    ]
    ws = wb.create_sheet("Notas")
    ws.column_dimensions["A"].width = 118
    for i, (t, neg) in enumerate(notas, start=1):
        c = ws.cell(i, 1, t)
        c.font = Font(name="Arial", size=10, bold=neg,
                      color="1F4E79" if neg else "1A1A1A")
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if not neg and t:
            ws.row_dimensions[i].height = 14 * (1 + len(t) // 108)

    for h in wb.sheetnames:
        wb[h].sheet_properties.tabColor = (
            "7F7F7F" if h in ("Notas", "Variables instrumentales", "Serie")
            else "1F4E79")
    out = os.path.join(SAL, "Elasticidad_sustitucion_AR_EC_Espana.xlsx")
    wb.save(out)
    print("Planilla:", out)


# ----------------------------------------------------------------------- main
def main() -> None:
    s, M = panel()
    print("=" * 98)
    print("ELASTICIDAD DE SUSTITUCIÓN LANGOSTINO ARGENTINO / VANNAMEI ECUATORIANO "
          "EN ESPAÑA")
    print("=" * 98)
    print(f"Mensual, {s.index[0]} a {s.index[-1]}, {len(s)} observaciones.")
    print(f"ln(q_ar/q_ec) va de {s.lq.iloc[:12].mean():+.2f} en 2013 a "
          f"{s.lq.iloc[-12:].mean():+.2f} en los últimos doce meses.")
    print(f"ln(p_ec/p_ar) va de {s.lp.iloc[:12].mean():+.2f} a "
          f"{s.lp.iloc[-12:].mean():+.2f}.")

    tabs = {}
    tabs["raices"] = raices(s)
    print("\n" + "-" * 98 + "\nRAÍCES UNITARIAS\n" + "-" * 98)
    print(tabs["raices"].round(3).to_string(index=False))

    tabs["coint"], beta = cointegracion(s)
    print("\n" + "-" * 98 + "\nCOINTEGRACIÓN Y CAUSALIDAD\n" + "-" * 98)
    print(tabs["coint"].round(4).to_string(index=False))

    tabs["est"], g = estimaciones(s, M)
    print("\n" + "-" * 98 + "\nESTIMACIONES DE LA ELASTICIDAD\n" + "-" * 98)
    print(tabs["est"].round(3).to_string(index=False))

    tabs["ecm"], _ = ecm(s, M, g["ols"])
    print("\n" + "-" * 98 + "\nCORTO PLAZO (corrección de error)\n" + "-" * 98)
    print(tabs["ecm"].round(3).to_string(index=False))

    sigma = float(g["dols"].params["lp"])
    tabs["desc"] = descomposicion(s, sigma)
    print("\n" + "-" * 98 + f"\nDESCOMPOSICIÓN con s = {sigma:.2f}\n" + "-" * 98)
    print(tabs["desc"].round(3).to_string(index=False))

    tabs["iv"] = iv(s, M)
    print("\n" + "-" * 98 + "\nVARIABLES INSTRUMENTALES (documentado, no se usa)\n"
          + "-" * 98)
    print(tabs["iv"].round(3).to_string(index=False))

    ci = g["dols"].conf_int().loc["lp"]
    print("\n" + "=" * 98)
    print(f"RESULTADO: elasticidad de sustitución de largo plazo = {sigma:.2f} "
          f"(IC95 {ci[0]:.2f} a {ci[1]:.2f})")
    print(f"           Johansen implica {beta:.2f}; el rango de las "
          f"especificaciones va de "
          f"{tabs['est']['Elasticidad s'].drop(2).min():.2f} a "
          f"{tabs['est']['Elasticidad s'].drop(2).max():.2f}")
    print("=" * 98)

    grafico(s, sigma, float(g["dols"].params["const"]))
    planilla(s, tabs)


if __name__ == "__main__":
    main()
