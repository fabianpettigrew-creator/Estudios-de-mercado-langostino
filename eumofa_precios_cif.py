# -*- coding: utf-8 -*-
"""Planilla de precios CIF de EUMOFA: langostino argentino contra vannamei ecuatoriano.

Qué es cada serie. EUMOFA publica el comercio declarado por los países de la UE
(Eurostat/Comext). Se toma la importación EXTRA-UE, congelada, y se calcula el valor
unitario mes a mes, que es valor declarado dividido volumen:

  · Langostino argentino  «Shrimp, miscellaneous», origen Argentina. Es la partida
    donde entra Pleoticus muelleri; el warmwater argentino existe pero es el 0,1% del
    volumen y queda afuera, igual que en el resto del estudio.
  · Vannamei ecuatoriano  «Shrimp, warmwater», origen Ecuador.

Los valores de importación de Comext son CIF en frontera de la UE, así que estos
precios NO son comparables con el FOB de despacho argentino: incluyen flete y seguro
hasta el puerto europeo. Sirven para comparar los dos orígenes entre sí, que es
exactamente para lo que se usan acá.

La hoja 2024-2026 se recalcula desde los CSV de EUMOFA y se contrasta contra
`datos/eumofa_camaron_ue.pkl`, que es el insumo que usa el modelo. La hoja de la serie
completa 2013-2026 sale de ese pickle.

Uso:  python eumofa_precios_cif.py
Salida: salidas/EUMOFA_precios_CIF_langostino_vannamei.xlsx
"""
from __future__ import annotations

import os
import sys

import numpy as np
import pandas as pd

sys.stdout.reconfigure(encoding="utf-8")
RAIZ = os.path.dirname(os.path.abspath(__file__))
EUMOFA = (r"C:\Users\fmpet\OneDrive\IA agentes\DATOS PESCA UE PROYECTO CLAUDE"
          r"\Bases para analizar")
SALIDA = os.path.join(RAIZ, "salidas",
                      "EUMOFA_precios_CIF_langostino_vannamei.xlsx")

COLS = ["year", "month", "flow_type", "intra_extra_EU", "partner_contry",
        "main_commercial_species", "presentation", "preservation",
        "value(EUR)", "volume(kg)"]
CONGELADO = ["PS2 Frozen", "PS5 Unspecified"]
MESES = ["enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto",
         "septiembre", "octubre", "noviembre", "diciembre"]


def _archivos(anios) -> list[str]:
    """2013-2025 están en la descarga masiva; el año corriente es el parcial suelto."""
    out = []
    for y in anios:
        masiva = os.path.join(EUMOFA, "EUMOFA BASES", "DESCARGA MASIVA", "EXPO IMPO",
                              "EU", f"{y}_Trade_data_reported_by_EU_countries.csv")
        suelto = os.path.join(EUMOFA, f"{y}_Trade_data_reported_by_EU_countries.csv")
        r = masiva if os.path.exists(masiva) else suelto
        if not os.path.exists(r):
            sys.exit(f"ABORTA: no está el archivo de EUMOFA de {y}")
        out.append(r)
    return out


def crudo(anios=(2024, 2025, 2026)) -> pd.DataFrame:
    tr = []
    for f in _archivos(anios):
        d = pd.read_csv(f, sep=";", dtype=str, low_memory=False,
                        usecols=lambda c: c in COLS)
        d = d[(d.flow_type == "Import") & (d.intra_extra_EU == "Extra EU")
              & (d.partner_contry.isin(["Argentina", "Ecuador"]))
              & (d.preservation.isin(CONGELADO))
              & (d.main_commercial_species.isin(["Shrimp, miscellaneous",
                                                 "Shrimp, warmwater"]))]
        tr.append(d)
        print(f"  {os.path.basename(f)}: {len(d)} filas")
    d = pd.concat(tr, ignore_index=True)
    d["eur"] = pd.to_numeric(d["value(EUR)"], errors="coerce")
    d["kg"] = pd.to_numeric(d["volume(kg)"], errors="coerce")
    d["per"] = pd.PeriodIndex(
        pd.to_numeric(d.year).astype(int).astype(str) + "-"
        + pd.to_numeric(d.month).astype(int).astype(str).str.zfill(2), freq="M")
    return d


def mensual(d: pd.DataFrame) -> pd.DataFrame:
    """Valor unitario mensual de cada origen, con su valor y su volumen."""
    def bloque(sub, pre):
        g = sub.groupby("per")[["eur", "kg"]].sum()
        return pd.DataFrame({f"{pre}_eur": g.eur, f"{pre}_t": g.kg / 1000,
                             f"{pre}_eur_kg": g.eur / g.kg})

    ar = bloque(d[(d.partner_contry == "Argentina")
                  & (d.main_commercial_species == "Shrimp, miscellaneous")], "ar")
    ec = bloque(d[(d.partner_contry == "Ecuador")
                  & (d.main_commercial_species == "Shrimp, warmwater")], "ec")
    m = ar.join(ec, how="outer").sort_index()
    m["dif_eur_kg"] = m.ar_eur_kg - m.ec_eur_kg
    m["dif_pct"] = (m.ar_eur_kg / m.ec_eur_kg - 1) * 100
    m["cociente"] = m.ar_eur_kg / m.ec_eur_kg
    return m


def verificar(m: pd.DataFrame) -> None:
    """El recálculo tiene que dar idéntico al pickle que consume el modelo."""
    P = pd.read_pickle(os.path.join(RAIZ, "datos", "eumofa_camaron_ue.pkl"))
    P = P.reindex(m.index)
    print("\n" + "-" * 88)
    print("CONTRASTE CONTRA datos/eumofa_camaron_ue.pkl (el insumo del modelo)")
    print("-" * 88)
    for a, b, nom in [("ar_eur_kg", "ar_eur_kg", "langostino argentino"),
                      ("ec_eur_kg", "van_ec_eur_kg", "vannamei ecuatoriano")]:
        dif = (m[a] - P[b]).abs().max()
        print(f"   {nom:<24s} máxima diferencia {dif:.2e} EUR/kg   "
              + ("idéntico" if dif < 1e-9 else "REVISAR"))


def anual(m: pd.DataFrame) -> pd.DataFrame:
    a = m.groupby(m.index.year).agg(
        ar_eur=("ar_eur", "sum"), ar_t=("ar_t", "sum"),
        ec_eur=("ec_eur", "sum"), ec_t=("ec_t", "sum"), meses=("ar_t", "size"))
    a["ar_eur_kg"] = a.ar_eur / (a.ar_t * 1000)
    a["ec_eur_kg"] = a.ec_eur / (a.ec_t * 1000)
    a["dif_eur_kg"] = a.ar_eur_kg - a.ec_eur_kg
    a["dif_pct"] = (a.ar_eur_kg / a.ec_eur_kg - 1) * 100
    a.index.name = "año"
    return a[["ar_eur", "ar_t", "ar_eur_kg", "ec_eur", "ec_t", "ec_eur_kg",
              "dif_eur_kg", "dif_pct", "meses"]]


# --------------------------------------------------------------------- planilla
ENC = {"per": "Mes", "anio": "Año", "mes": "Mes",
       "ar_eur": "Langostino AR · valor CIF (EUR)",
       "ar_t": "Langostino AR · volumen (t)",
       "ar_eur_kg": "Langostino AR · precio CIF (EUR/kg)",
       "ec_eur": "Vannamei EC · valor CIF (EUR)",
       "ec_t": "Vannamei EC · volumen (t)",
       "ec_eur_kg": "Vannamei EC · precio CIF (EUR/kg)",
       "dif_eur_kg": "Diferencia AR − EC (EUR/kg)",
       "dif_pct": "Diferencia AR / EC (%)",
       "cociente": "Cociente AR / EC",
       "meses": "Meses con dato"}

NOTAS = [
    ("Qué contiene esta planilla", True),
    ("Precios mensuales de importación de la Unión Europea, en euros por kilo, del "
     "langostino argentino y del camarón vannamei ecuatoriano. Fuente: EUMOFA, "
     "«Trade data reported by EU countries» (Eurostat / Comext).", False),
    ("", False),
    ("Definición de cada serie", True),
    ("Langostino argentino: importación extra-UE, congelada, especie comercial "
     "«Shrimp, miscellaneous», origen Argentina. Es la partida donde entra Pleoticus "
     "muelleri. El «Shrimp, warmwater» de origen argentino existe pero es el 0,1% del "
     "volumen y queda afuera, igual que en el resto del estudio.", False),
    ("Vannamei ecuatoriano: importación extra-UE, congelada, especie comercial "
     "«Shrimp, warmwater», origen Ecuador.", False),
    ("Congelado incluye las categorías PS2 Frozen y PS5 Unspecified. Queda afuera "
     "PS4 Prepared/Preserved, que es producto elaborado y cotiza muy por encima.", False),
    ("", False),
    ("El precio es un valor unitario", True),
    ("No es una cotización de mercado: es el valor declarado dividido por el volumen "
     "declarado, mes a mes. Se mueve con el precio y también con la mezcla de "
     "producto y de calibre que se embarcó ese mes.", False),
    ("", False),
    ("CIF, no FOB", True),
    ("Los valores de importación de Comext son CIF en frontera de la UE: incluyen "
     "flete y seguro hasta el puerto europeo. NO son comparables con el FOB de "
     "despacho argentino que usa el estudio para el precio del L1 tangonero. Sirven "
     "para comparar los dos orígenes entre sí, que es para lo que están acá.", False),
    ("", False),
    ("Sin apertura por presentación", True),
    ("A este nivel de agregación EUMOFA declara todo el congelado como PR1 "
     "Whole/Gutted, así que no se puede separar entero de cola ni de pelado. Esa "
     "apertura está en el archivo CN8_details, que no se procesó para esta planilla.", False),
    ("", False),
    ("Cobertura", True),
    ("La hoja mensual va de enero de 2024 a mayo de 2026: es hasta donde llega el "
     "archivo parcial de 2026 descargado. La hoja de la serie completa arranca en "
     "enero de 2013 y sale de datos/eumofa_camaron_ue.pkl, el mismo insumo que "
     "consume el modelo de demanda inversa; el recálculo de 2024-2026 se contrastó "
     "contra ese pickle y da idéntico.", False),
    ("", False),
    ("Reproducibilidad", True),
    ("«eumofa_precios_cif.py» en la raíz del proyecto. Elaboración: Lic. Fabián "
     "Pettigrew · AXIA.", False),
]


def planilla(m: pd.DataFrame, an: pd.DataFrame, larga: pd.DataFrame) -> None:
    from openpyxl.chart import LineChart, Reference
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    AZUL, GRIS = "1F4E79", "F2F6F8"
    tit = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    relleno = PatternFill("solid", fgColor=AZUL)
    borde = Border(bottom=Side(style="thin", color="BFBFBF"))

    def volcar(ws, df, indice):
        cols = [indice] + list(df.columns)
        for j, c in enumerate(cols, start=1):
            cel = ws.cell(row=1, column=j, value=ENC.get(c, c))
            cel.font, cel.fill = tit, relleno
            cel.alignment = Alignment(wrap_text=True, vertical="center",
                                      horizontal="center")
        ws.row_dimensions[1].height = 34
        for i, (k, fila) in enumerate(df.iterrows(), start=2):
            ws.cell(row=i, column=1, value=str(k)).border = borde
            for j, c in enumerate(df.columns, start=2):
                v = fila[c]
                cel = ws.cell(row=i, column=j,
                              value=None if pd.isna(v) else float(v))
                cel.border = borde
                cel.number_format = ("#,##0" if c.endswith(("_eur", "_t", "meses"))
                                     else "#,##0.000" if c.endswith("_eur_kg")
                                     or c in ("dif_eur_kg", "cociente")
                                     else '+#,##0.0"%";-#,##0.0"%"')
        ws.column_dimensions["A"].width = 11
        for j in range(2, len(cols) + 1):
            ws.column_dimensions[get_column_letter(j)].width = 17
        ws.freeze_panes = "B2"

    with pd.ExcelWriter(SALIDA, engine="openpyxl") as w:
        libro = w.book
        libro.remove(libro.active) if libro.sheetnames else None

        h1 = libro.create_sheet("2024-2026 mensual")
        volcar(h1, m, "per")
        h2 = libro.create_sheet("2024-2026 anual")
        volcar(h2, an, "anio")
        h3 = libro.create_sheet("Serie completa 2013-2026")
        volcar(h3, larga, "per")

        # gráfico nativo sobre la hoja mensual
        g = LineChart()
        g.title = "Precio CIF de importación en la UE, EUR/kg"
        g.height, g.width = 9, 22
        g.y_axis.title = "EUR por kilo"
        for col in (4, 7):      # ar_eur_kg y ec_eur_kg
            g.add_data(Reference(h1, min_col=col, min_row=1, max_row=len(m) + 1),
                       titles_from_data=True)
        g.set_categories(Reference(h1, min_col=1, min_row=2, max_row=len(m) + 1))
        g.series[0].graphicalProperties.line.solidFill = "1F4E79"
        g.series[1].graphicalProperties.line.solidFill = "C0504D"
        for s in g.series:
            s.graphicalProperties.line.width = 22000
            s.smooth = False
        h1.add_chart(g, "M2")

        h4 = libro.create_sheet("Notas")
        h4.column_dimensions["A"].width = 118
        for i, (t, neg) in enumerate(NOTAS, start=1):
            cel = h4.cell(row=i, column=1, value=t)
            cel.font = Font(name="Calibri", size=11, bold=neg,
                            color=AZUL if neg else "1A1A1A")
            cel.alignment = Alignment(wrap_text=True, vertical="top")
            if not neg and t:
                h4.row_dimensions[i].height = 15 * (1 + len(t) // 110)
        for h in (h1, h2, h3):
            h.sheet_properties.tabColor = AZUL
        h4.sheet_properties.tabColor = "7F7F7F"


def main() -> None:
    print("=" * 88)
    print("PRECIOS CIF DE EUMOFA · LANGOSTINO ARGENTINO vs VANNAMEI ECUATORIANO")
    print("=" * 88)
    d = crudo()
    m = mensual(d)
    verificar(m)
    an = anual(m)

    P = pd.read_pickle(os.path.join(RAIZ, "datos", "eumofa_camaron_ue.pkl"))
    larga = pd.DataFrame({"ar_eur_kg": P.ar_eur_kg, "ar_t": P.ar_kt * 1000,
                          "ec_eur_kg": P.van_ec_eur_kg, "ec_t": P.van_ec_kt * 1000})
    larga["dif_eur_kg"] = larga.ar_eur_kg - larga.ec_eur_kg
    larga["dif_pct"] = (larga.ar_eur_kg / larga.ec_eur_kg - 1) * 100
    larga["cociente"] = larga.ar_eur_kg / larga.ec_eur_kg

    print("\n" + "-" * 88)
    print("MENSUAL 2024-2026 · precio CIF en EUR/kg")
    print("-" * 88)
    print(m[["ar_eur_kg", "ar_t", "ec_eur_kg", "ec_t", "dif_eur_kg",
             "dif_pct"]].round(3).to_string())
    print("\n" + "-" * 88)
    print("ANUAL · valor unitario ponderado por volumen")
    print("-" * 88)
    print(an[["ar_eur_kg", "ar_t", "ec_eur_kg", "ec_t", "dif_eur_kg", "dif_pct",
              "meses"]].round(2).to_string())

    try:
        planilla(m, an, larga)
        print(f"\nGuardado: {SALIDA}")
    except PermissionError:
        print(f"\nAVISO: no pude escribir {SALIDA} (¿abierto en Excel?)")


if __name__ == "__main__":
    main()
